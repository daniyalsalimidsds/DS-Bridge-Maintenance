#!/usr/bin/env python3
"""Extract the supplied inspection form, preserving exact names, weights and provenance.

Read-only XLSX extraction; the original workbook is never modified. Mapping decisions
are explicit and reviewed, not based on fuzzy name matching at application runtime.
"""
import argparse
import hashlib
import json
from pathlib import Path
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]

# row number: existing semantic equivalents. Lists keep split defects intact.
MATCH = {
 1: ['road_pavement-1','road_pavement-2','road_pavement-3','road_pavement-6','road_pavement-7','road_pavement-8','road_pavement-9'],
 3: ['road_pavement-5'], 5: ['approach-1'], 7: ['barriers-6'],
 14: ['barriers-4'], 15: ['barriers-1','barriers-2'], 16: ['barriers-5'],
 17: ['barriers-3'], 22: ['drainage-1','drainage-8'],
 25: ['drainage-2','drainage-3'], 28: ['joints-1','joints-3','joints-8'],
 29: ['joints-7'], 30: ['joints-4'], 31: ['joints-6'], 33: ['joints-2'],
 35: ['utilities-1'], 38: ['utilities-8'], 39: ['utilities-3'], 40: ['utilities-2'],
 46: ['alignment-7'], 47: ['alignment-2'],
 56: ['bearings-1','bearings-2','bearings-3','bearings-6','bearings-7','bearings-8','bearings-10'] + ['neoprene-'+str(i) for i in list(range(1,10))+list(range(11,17))+[18]],
 57: ['bearings-4','bearings-5','neoprene-10'],
 60: ['bearings-9','neoprene-17'], 64: ['piers-4','abutments-2'],
 65: ['foundation-5'], 66: ['hydraulic-1','piers-7','foundation-1','foundation-6'],
 67: ['foundation-9','hydraulic-7'], 69: ['piers-9'],
 70: ['abutments-4','abutments-5','approach-5'],
}
# Only one-to-one equivalences are renamed. Combined or more detailed originals stay.
RENAME = {3:'road_pavement-5',5:'approach-1',14:'barriers-4',16:'barriers-5',
          29:'joints-7',30:'joints-4',31:'joints-6',33:'joints-2',38:'utilities-8',
          39:'utilities-3',46:'alignment-7',47:'alignment-2',
          65:'foundation-5',69:'piers-9'}

def material(row, material_id, indices, scopes):
    MATCH[row] = [f'v12-{material_id}-{i}--{scope}' for scope in scopes for i in indices]

deck = ['05_عرشه_ابرسازه','06_کف_عرشه_دال']
sub = ['07_پایه_میانی','08_کوله']
for scopes, offset in [(deck,0),(sub,23)]:
    material(48+offset,'concrete',list(range(11,26))+[37,38,39,42,43],scopes)
    material(49+offset,'concrete',[6,7,29,33,34,35,36],scopes)
    material(50+offset,'concrete',[1,2,3,4,5,8,40],scopes)
    material(51+offset,'concrete',[9,10,26,44,45,46],scopes)
    material(52+offset,'steel',list(range(1,9))+[23,50],scopes)
    material(53+offset,'steel',[20,21,22,24,25,26,27],scopes)
    material(54+offset,'steel',[28,29,30,31,32,33,34],scopes)
    material(55+offset,'steel',[13,14,15,16,36],scopes)
material(79,'concrete',[27,28],deck)

# Added rows belong to the relevant existing group. Presence checks are supplied
# separately, since 'absence of a required safety feature' remains inspectable.
GROUP = {}
for numbers, group in [
    (range(1,6),'road_safety'),(range(6,14),'barriers'),(range(14,21),'barriers'),
    (range(21,28),'drainage'),(range(28,34),'joints'),(range(34,42),'utilities'),
    (range(42,48),'deck'),(range(56,61),'bearings'),(range(61,71),'piers')]:
    for n in numbers: GROUP[n] = group

def build(path):
    book = load_workbook(path, data_only=False)
    sheet = book['Ins. Form']
    rows=[]; location=''
    for r in list(range(22,63))+list(range(73,111)):
        n = sheet.cell(r,20).value
        if not isinstance(n,int): continue
        if sheet.cell(r,19).value: location=sheet.cell(r,19).value
        title=sheet.cell(r,13).value
        if not title: continue
        rows.append(dict(number=n, id=f'M{n:02}', name=title, weight=sheet.cell(r,1).value,
                         location=location, sourceCell=f"'Ins. Form'!M{r}",weightCell=f"'Ins. Form'!A{r}",
                         existing=MATCH.get(n,[]), rename=RENAME.get(n),
                         addTo=GROUP.get(n,'deck') if n not in MATCH else None,
                         customInSource=n==79, sampleStatus=next((key for key,col in [('uninspectable',8),('none',9),('low',10),('medium',11),('emergency',12)] if sheet.cell(r,col).value),None)))
    assert len(rows)==79 and all(1<=r['weight']<=5 for r in rows)
    scales = dict(zip(['uninspectable','none','low','medium','emergency'],[sheet.cell(21,c).value for c in range(8,13)]))
    assert scales=={'uninspectable':-10,'none':0,'low':-10,'medium':-30,'emergency':-100}
    sample_total=sum(scales[row['sampleStatus']]*row['weight'] for row in rows if row['sampleStatus'])
    assert sample_total==-240
    result={'version':'1.6.0','sourceFile':path.name,'sourceSha256':hashlib.sha256(path.read_bytes()).hexdigest(),
            'sourceSheet':'Ins. Form','sourceDocument':'621-8-6 — عنوان درج‌شده در فایل ارسالی؛ نسخه مستقل آیین‌نامه احراز نشده',
            'scores':scales,'rows':rows,
            'verification':{'low':-80,'uninspectable':-160,'total':-240,'populatedRows':79,
                            'note':'A118=240 قدر مطلق نمره منفی است؛ A115=302 جمع ضرایب شامل چهار ردیف خالی است و شاخص سلامت نیست.'}}
    output=ROOT/'app/src/main/assets/data/municipal-catalog.js'
    output.write_text('/* Generated by tools/import_municipal_catalog.py. */\nwindow.BRIDGE_MUNICIPAL_CATALOG = '+json.dumps(result,ensure_ascii=False,indent=2)+';\n')
    (ROOT/'reference/municipal-catalog.json').write_text(json.dumps(result,ensure_ascii=False,indent=2))
    print(f'Extracted {len(rows)} populated damage rows; {sum(not x["existing"] for x in rows)} additions; {len(RENAME)} exact-equivalent renames.')

if __name__=='__main__':
    p=argparse.ArgumentParser();p.add_argument('workbook',type=Path);build(p.parse_args().workbook)
