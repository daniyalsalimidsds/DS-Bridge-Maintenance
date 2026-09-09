#!/usr/bin/env python3
"""Generate the in-app bridge profile schema from the Publication 367 workbook."""

from __future__ import annotations

import argparse
import hashlib
import json
import re
from pathlib import Path

import openpyxl
from openpyxl.utils import range_boundaries


def slug(value: str) -> str:
    value = re.sub(r"[^0-9A-Za-z_\u0600-\u06ff]+", "_", value.strip())
    return value.strip("_") or "field"


def build_validation_index(sheet):
    result = {}
    for validation in sheet.data_validations.dataValidation:
        if validation.type != "list" or not validation.formula1:
            continue
        source = validation.formula1.lstrip("=")
        source = source.split("!", 1)[-1].replace("'", "")
        min_col, min_row, max_col, max_row = range_boundaries(source)
        options = []
        for row in sheet.iter_rows(
            min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col
        ):
            for cell in row:
                if cell.value not in (None, ""):
                    options.append(str(cell.value).strip())
        for cell_range in validation.ranges.ranges:
            selected = sheet[cell_range.coord]
            if not isinstance(selected, tuple):
                selected = ((selected,),)
            elif selected and not isinstance(selected[0], tuple):
                selected = (selected,)
            for row in selected:
                for cell in row:
                    result[cell.coordinate] = options
    return result


def infer_kind(label: str, options):
    if options:
        return "select"
    if any(word in label for word in ("تعداد", "طول کل")):
        return "number"
    return "text"


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("workbook", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()

    workbook = openpyxl.load_workbook(args.workbook, data_only=False)
    sheet = workbook["01_خلاصه_پروفایل_منویی"]
    validations = build_validation_index(sheet)
    sections = []

    for row_number in range(8, 28):
        title = sheet.cell(row_number, 1).value
        code = sheet.cell(row_number, 2).value
        if not title or not code:
            continue
        fields = []
        for index, (label_col, value_col) in enumerate(((3, 4), (5, 6), (7, 8)), 1):
            label = sheet.cell(row_number, label_col).value
            if not label:
                continue
            value_cell = sheet.cell(row_number, value_col)
            options = validations.get(value_cell.coordinate)
            field_id = f"{code}__{index}_{slug(str(label))}"
            fields.append(
                {
                    "id": field_id,
                    "label": str(label).strip(),
                    "kind": infer_kind(str(label), options),
                    "options": options or [],
                    "default": "",
                    "example": value_cell.value if value_cell.value is not None else "",
                    "required": str(label).strip() in {"نام پل", "کاربری اصلی"},
                }
            )
        sections.append({"id": str(code), "title": str(title).strip(), "fields": fields})

    source_bytes = args.workbook.read_bytes()
    payload = {
        "version": "367-menu-profile-v1",
        "source": args.workbook.name,
        "sourceSha256": hashlib.sha256(source_bytes).hexdigest(),
        "sections": sections,
        "locationFields": [
            {"id": "latitude", "label": "عرض جغرافیایی", "kind": "number", "required": False},
            {"id": "longitude", "label": "طول جغرافیایی", "kind": "number", "required": False},
            {"id": "accuracyM", "label": "دقت GPS (متر)", "kind": "number", "required": False},
            {"id": "capturedAt", "label": "زمان ثبت موقعیت", "kind": "datetime", "required": False},
            {"id": "source", "label": "روش ثبت موقعیت", "kind": "select", "options": ["GPS", "دستی"], "required": False},
        ],
    }

    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(
        "window.BRIDGE_PROFILE_SCHEMA = Object.freeze("
        + json.dumps(payload, ensure_ascii=False, separators=(",", ":"))
        + ");\n",
        encoding="utf-8",
    )


if __name__ == "__main__":
    main()
